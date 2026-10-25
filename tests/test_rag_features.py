"""核心能力单测：断崖截断、上下文预算、引用后置校验、正文清洗、Office 解析。

全部不依赖网络 / Milvus / API key，可在 CI 里直接跑。
"""
import math

import pytest

from app.query_process.agent.nodes.node_rerank import select_by_cliff
from app.utils.citation_utils import extract_citations, validate_and_fix
from app.utils.text_clean_utils import sanitize_content
from app.utils.token_budget_utils import build_context, estimate_tokens


# ==================== 断崖检测动态 Top-K ====================
def test_cliff_empty():
    assert select_by_cliff([]) == []


def test_cliff_keeps_all_when_scores_are_close():
    ranked = [(0, 0.90), (1, 0.88), (2, 0.85)]
    assert select_by_cliff(ranked) == [0, 1, 2]


def test_cliff_cuts_at_sharp_drop():
    # 0.30 / 0.85 = 0.35 < 0.6，应在第 3 条前截断
    ranked = [(0, 0.90), (1, 0.85), (2, 0.30), (3, 0.28)]
    assert select_by_cliff(ranked) == [0, 1]


def test_cliff_falls_back_to_two_when_all_below_threshold():
    # 全部低于 0.30，但 top1 高于软下限 0.15：保底给 2 条，交给 prompt 兜底
    ranked = [(0, 0.29), (1, 0.28)]
    assert select_by_cliff(ranked) == [0, 1]


def test_cliff_returns_empty_when_top1_is_too_low():
    # 连软下限都不到，直接判定无相关内容
    ranked = [(0, 0.10), (1, 0.05)]
    assert select_by_cliff(ranked) == []


def test_cliff_fallback_when_below_threshold_but_above_soft_floor():
    # 全部低于 0.30，但 top1 高于软下限 0.15：保底给 2 条
    ranked = [(0, 0.20), (1, 0.18), (2, 0.17)]
    assert select_by_cliff(ranked) == [0, 1]


def test_cliff_respects_max_keep():
    ranked = [(index, 0.9) for index in range(20)]
    assert len(select_by_cliff(ranked, max_keep=8)) == 8


def test_cliff_keeps_original_indexes():
    ranked = [(5, 0.9), (2, 0.8), (9, 0.7)]
    assert select_by_cliff(ranked) == [5, 2, 9]


# ==================== 上下文预算 ====================
def test_estimate_tokens_prefers_cjk():
    assert estimate_tokens("") == 0
    # 中文按 1 字 1 token 估
    assert estimate_tokens("打印机") == 3


def test_build_context_uses_all_when_budget_is_large():
    docs = [{"content": "第一段"}, {"content": "第二段"}]
    context, used, info = build_context(docs, budget_tokens=1000)
    assert len(used) == 2
    assert "[1] 第一段" in context and "[2] 第二段" in context
    assert info["truncated"] is False


def test_build_context_drops_when_budget_is_small():
    docs = [{"content": "很长的内容" * 100}, {"content": "第二段"}]
    _context, used, info = build_context(docs, budget_tokens=50)
    assert used == []
    assert info["truncated"] is True
    assert info["dropped_count"] == 2


def test_build_context_skips_empty_content():
    docs = [{"content": ""}, {"content": "有效内容"}]
    _context, used, info = build_context(docs, budget_tokens=100)
    assert len(used) == 1
    assert info["dropped_count"] == 1


def test_build_context_numbering_matches_references():
    docs = [{"content": f"内容{i}"} for i in range(3)]
    context, _used, _info = build_context(docs, budget_tokens=1000)
    for index in range(3):
        assert f"[{index + 1}] 内容{index}" in context


# ==================== 引用后置校验 ====================
def test_extract_citations():
    assert extract_citations("见 [1]，另见[2]") == [1, 2]
    assert extract_citations("没有引用") == []
    # 年份不应被当成引用
    assert extract_citations("2024 年发布") == []


def test_validate_removes_out_of_range_citations():
    answer = "温度为 120 度[1]，参见[7]。"
    fixed, info = validate_and_fix(answer, reference_count=2)
    assert "[7]" not in fixed
    assert "[1]" in fixed
    assert info["invalid_citations"] == [7]
    assert info["has_invalid"] is True


def test_validate_keeps_all_valid_citations():
    answer = "第一步[1]，第二步[2]。"
    fixed, info = validate_and_fix(answer, reference_count=2)
    assert fixed == answer
    assert info["has_invalid"] is False
    assert info["all_referenced"] is True


def test_validate_with_zero_references_removes_everything():
    fixed, info = validate_and_fix("随便说点什么[1]", reference_count=0)
    assert "[1]" not in fixed
    assert info["invalid_citations"] == [1]


# ==================== 正文清洗 ====================
def test_sanitize_replaces_legacy_minio_image():
    text = "见图 ![打印机面板](http://localhost:9000/knowledge-base-files/a.jpg) 说明"
    cleaned = sanitize_content(text)
    assert "localhost:9000" not in cleaned
    assert "[图片：打印机面板]" in cleaned


def test_sanitize_keeps_normal_markdown():
    text = "正常内容，带 [链接](https://example.com/doc)"
    assert sanitize_content(text) == text


def test_sanitize_handles_empty_alt():
    text = "![](http://127.0.0.1:9000/x.png)"
    assert sanitize_content(text) == "[图片]"


# ==================== Office 解析 ====================
def test_docx_to_markdown(tmp_path):
    from docx import Document

    from app.import_process.agent.nodes.node_office_to_md import docx_to_markdown

    document = Document()
    document.add_heading("产品概述", level=1)
    document.add_paragraph("这是一段正文。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "参数"
    table.cell(0, 1).text = "值"
    table.cell(1, 0).text = "电源"
    table.cell(1, 1).text = "220V"

    path = tmp_path / "manual.docx"
    document.save(path)

    markdown = docx_to_markdown(str(path))
    assert "# 产品概述" in markdown
    assert "这是一段正文。" in markdown
    assert "| 参数 | 值 |" in markdown
    assert "| 电源 | 220V |" in markdown


def test_xlsx_to_markdown_keeps_sheet_and_row_no(tmp_path):
    from openpyxl import Workbook

    from app.import_process.agent.nodes.node_office_to_md import xlsx_to_markdown

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "设备清单"
    sheet.append(["型号", "数量"])
    sheet.append(["ER2100", 3])

    path = tmp_path / "inventory.xlsx"
    workbook.save(path)

    markdown = xlsx_to_markdown(str(path))
    assert "## Sheet: 设备清单" in markdown
    assert "| 行号 |" in markdown
    assert "ER2100" in markdown


def test_office_node_rejects_unsupported_suffix(tmp_path):
    from app.import_process.agent.nodes.node_office_to_md import node_office_to_md

    path = tmp_path / "note.txt"
    path.write_text("hello", encoding="utf-8")

    state = {"task_id": "t1", "local_file_path": str(path), "local_dir": str(tmp_path)}
    with pytest.raises(ValueError):
        node_office_to_md(state)
    assert state["failed_stage"] == "office"
